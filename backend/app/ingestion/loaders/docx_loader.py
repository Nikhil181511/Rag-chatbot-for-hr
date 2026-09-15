import os
from typing import Optional, List, Dict, Any
from docx import Document as DocxDocument
from app.ingestion.loaders.base import DocumentLoader, NormalizedDocument, PageReference
from app.config.logging import get_logger

logger = get_logger(__name__)


class DOCXLoader(DocumentLoader):
    def load(self, file_path: str, original_filename: Optional[str] = None) -> NormalizedDocument:
        file_name = original_filename or os.path.basename(file_path)
        doc = DocxDocument(file_path)

        content_parts: List[str] = []
        sections: List[Dict[str, Any]] = []
        current_section = "General"

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            # Check heading style
            if paragraph.style.name.startswith("Heading"):
                current_section = text
                content_parts.append(f"\n## {text}\n")
                sections.append({"title": text, "level": paragraph.style.name})
            else:
                content_parts.append(text)

        # Extract tables
        tables_extracted: List[Dict[str, Any]] = []
        for t_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_rows.append(row_data)

            if table_rows and len(table_rows) > 1:
                headers = table_rows[0]
                rows = table_rows[1:]
                md_table = "\n| " + " | ".join(headers) + " |\n"
                md_table += "| " + " | ".join(["---"] * len(headers)) + " |\n"
                for row in rows:
                    md_table += "| " + " | ".join(row) + " |\n"
                content_parts.append(md_table)
                tables_extracted.append({"index": t_idx, "markdown": md_table})

        combined_content = "\n\n".join(content_parts)
        title = os.path.splitext(file_name)[0].replace("-", " ").replace("_", " ").title()

        return NormalizedDocument(
            source_file_name=file_name,
            source_type="docx",
            title=title,
            content=combined_content,
            sections=sections,
            tables=tables_extracted,
            page_references=[PageReference(page_number=1, text=combined_content)],
            metadata={"paragraphs_count": len(doc.paragraphs), "tables_count": len(doc.tables)},
        )
